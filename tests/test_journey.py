import importlib.util
import sys
import unittest
from datetime import date
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
def module(name, relative):
    spec = importlib.util.spec_from_file_location(name, ROOT / relative)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result
common = module('common', 'skills/yoga-journey-tracker/scripts/collect/_common.py')
aggregate = module('aggregate', 'skills/yoga-journey-tracker/scripts/analyze/aggregate_stats.py')

class JourneyTests(unittest.TestCase):
    def setUp(self):
        self.ws = openpyxl.Workbook().active
        self.ws.append(common.COLUMNS)
    def row(self, **changes):
        row = {'Date': date(2025, 1, 1), 'Time': '8:00 AM', 'Studio': 'Example Studio', 'Class Name (Raw)': 'Flow', 'Teacher': 'Teacher A', 'Source': 'Synthetic', 'Check-in Confirmed': 'Y'}
        row.update(changes)
        return row
    def test_multiple_sessions_and_reimport(self):
        self.assertEqual(common.upsert_visit(self.ws, self.row()), 'inserted')
        self.assertEqual(common.upsert_visit(self.ws, self.row(Time='6:00 PM')), 'inserted')
        self.assertEqual(common.upsert_visit(self.ws, self.row(Time='08:00')), 'stamped')
        self.assertEqual(self.ws.max_row, 3)
    def test_missing_time_is_not_wildcard(self):
        common.upsert_visit(self.ws, self.row())
        common.upsert_visit(self.ws, self.row(Time=''))
        self.assertEqual(self.ws.max_row, 3)
        self.assertEqual(common.upsert_visit(self.ws, self.row(Time='')), 'skipped')
    def test_different_class_or_city_does_not_merge(self):
        common.upsert_visit(self.ws, self.row(City='City A'))
        common.upsert_visit(self.ws, self.row(City='City B'))
        common.upsert_visit(self.ws, self.row(**{'Class Name (Raw)': 'Yin'}))
        self.assertEqual(self.ws.max_row, 4)
    def test_counts_attendance_days_and_individual_teachers(self):
        for row in [self.row(Teacher='Teacher A + Teacher B'), self.row(Time='6:00 PM'), self.row(Time='10:00 AM', **{'Check-in Confirmed': '', 'Unsure Attended': 'Y'})]:
            common.append_row(self.ws, row)
        stats = aggregate.summarize(aggregate.load_rows(self.ws))
        self.assertEqual(stats['total_visits'], 2)
        self.assertEqual(stats['practice_days'], 1)
        self.assertEqual(stats['unique_teachers'], 2)
        self.assertEqual(stats['attendance_counts']['provisional'], 1)
    def test_cancellation_overrides_checkin_and_empty_is_provisional(self):
        self.assertEqual(aggregate.attendance({'Attendance Status':'Late Cancel','Check-in Confirmed':'Y'}), 'excluded')
        self.assertEqual(aggregate.attendance({}), 'provisional')
    def test_empty_and_hours_are_not_inferred(self):
        stats = aggregate.summarize([])
        self.assertEqual(stats['total_visits'], 0)
        self.assertIsNone(stats['first_date'])
        common.append_row(self.ws, self.row(Training='200hr YTT'))
        stats = aggregate.summarize(aggregate.load_rows(self.ws))
        self.assertEqual(stats['training_hours']['completed_formal']['entries_with_hours'], 0)
    def test_explicit_completed_hours_only(self):
        stats = aggregate.summarize([], {'featured_trainings':[{'status':'completed','hours':200},{'status':'planned','hours':300}], 'workshops':[{'name':'Unknown length'}]})
        self.assertEqual(stats['training_hours']['completed_formal']['known_hours'], 200)
        self.assertEqual(stats['training_hours']['workshops']['entries_without_hours'], 1)

if __name__ == '__main__':
    unittest.main()
